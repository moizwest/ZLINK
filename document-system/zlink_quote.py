"""Z-Link QUOTATION template — PERMANENT (approved by the MD 29 Oct 2026).

Quotations are the one exception to the banner/watermark design in zlink_doc.py.
Every Z-Link quotation STARTS with this typeset letterhead, measured off the
MD's reference image artwork/quotation-header-reference.png:

  line 1  Z-LINK TECHNICAL SERVICES CONTRACTING CO. LLC   Calibri Bold 17 pt, blue #2E74B5, left
  line 2  OFFICE 201, BAY SQUARE, BUSINESS BAY, DUBAI, UNITED ARAB EMIRATES
                                                          Calibri Bold 12 pt, black, left
  line 3  "TRN:" Calibri Bold + 104180458200003 Regular   11.5 pt, black, left
  line 4  red phone 055-9495833 | green envelope INFO@ZLINK.AE | blue globe WWW.ZLINK.AE
                                                          Calibri 12 pt, black, left
  rule    grey 3-D rule (#A6A6A6, 1.5 pt emboss) across the full text width
  title   QUOTATION                                       Calibri Bold 12.5 pt, blue #2E74B5, RIGHT

Body (clauses Q3-Q9 in TEMPLATE-SPEC.md): To/Ref two-column block, Dear ...,
price schedule with blue header row, Sub Total / VAT 5% / Grand Total,
amount in words, optional items, basis, exclusions, terms (incl. RAK Bank),
closing and signatory. A4, 15 mm margins, footer "<ref> | Page X of Y".

Format is changed here once, never per document.
"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x2E, 0x74, 0xB5)
BLUE_HEX = '2E74B5'
LIGHT_HEX = 'DEEAF6'
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x80, 0x80, 0x80)
RED = RGBColor(0xFF, 0x00, 0x00)
GREEN = RGBColor(0x00, 0xB0, 0x50)
FONT = 'Calibri'
BODY_PT = 9.5
MARGIN = 15.0          # mm, all sides
USABLE = 210.0 - 2 * MARGIN   # 180 mm

COMPANY = 'Z-LINK TECHNICAL SERVICES CONTRACTING CO. LLC'
ADDRESS = 'OFFICE 201, BAY SQUARE, BUSINESS BAY, DUBAI, UNITED ARAB EMIRATES'
TRN = '104180458200003'
PHONE, EMAIL, WEB = '055-9495833', 'INFO@ZLINK.AE', 'WWW.ZLINK.AE'


def _run(p, text, size=BODY_PT, bold=False, color=None, italic=False, font=FONT):
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def _spacing(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        from docx.enum.text import WD_LINE_SPACING
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)


def _bottom_rule(p, style='threeDEmboss', sz=12, color='A6A6A6', space=6):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    for k, v in (('w:val', style), ('w:sz', str(sz)), ('w:space', str(space)), ('w:color', color)):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def quotation_letterhead(doc):
    """Clause Q1 + Q2 — the approved quotation letterhead and title. Call first."""
    p = doc.add_paragraph(); _spacing(p, line=20)
    _run(p, COMPANY, 17, True, BLUE)
    p = doc.add_paragraph(); _spacing(p, line=15)
    _run(p, ADDRESS, 12, True, BLACK)
    p = doc.add_paragraph(); _spacing(p, line=14)
    _run(p, 'TRN: ', 11.5, True, BLACK); _run(p, TRN, 11.5, False, BLACK)
    p = doc.add_paragraph(); _spacing(p, line=15)
    _run(p, '\u260E ', 11, False, RED); _run(p, PHONE + ' | ', 12)
    _run(p, '\u2709 ', 11, False, GREEN); _run(p, EMAIL + ' | ', 12)
    _run(p, '\U0001F310 ', 11, False, BLUE); _run(p, WEB, 12)
    _bottom_rule(p)
    p = doc.add_paragraph(); _spacing(p, before=6, after=6)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(p, 'QUOTATION', 12.5, True, BLUE)
    return doc


def _footer(section, ref):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, f'{ref}  |  Page ', 7.5, color=GREY)
    for instr in ('PAGE', None, 'NUMPAGES'):
        if instr is None:
            _run(p, ' of ', 7.5, color=GREY); continue
        r = _run(p, '', 7.5, color=GREY)
        for t, v in (('begin', None), ('instr', instr), ('end', None)):
            if t == 'instr':
                e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = v
            else:
                e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t)
            r._r.append(e)


def new_quotation(ref):
    """A4, 15 mm margins, Calibri 9.5 pt, footer with reference and Page X of Y,
    letterhead and title already placed. Returns a python-docx Document."""
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.paragraph_format.space_after = Pt(0)
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(MARGIN)
    s.header_distance, s.footer_distance = Mm(8), Mm(7)
    _footer(s, ref)
    quotation_letterhead(doc)
    return doc


# ---------- body helpers (clauses Q3-Q9) ----------
def para(doc, text='', size=BODY_PT, bold=False, color=None, align=None, before=0, after=2, italic=False):
    p = doc.add_paragraph(); _spacing(p, before, after)
    if align is not None:
        p.alignment = align
    if text:
        _run(p, text, size, bold, color, italic)
    return p


def heading(doc, text):
    """Section heading: blue bold 11 pt, kept with next."""
    p = para(doc, text, 11, True, BLUE, before=8, after=3)
    p.paragraph_format.keep_with_next = True
    return p


def set_widths(table, widths_mm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths_mm):
        gc.set(qn('w:w'), str(int(w * 56.6929)))
    for row in table.rows:
        for c, w in zip(row.cells, widths_mm):
            c.width = Mm(w)


def borders(table, color='999999', sz=4, none=False):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'nil' if none else 'single')
        if not none:
            e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_fill)
    tcPr.append(sh)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); trPr.append(e)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); trPr.append(e)


def cell_text(cell, runs, align=None, valign=WD_ALIGN_VERTICAL.CENTER):
    """runs: list of paragraphs; each paragraph a list of (text, bold[, color[, italic]])."""
    cell.vertical_alignment = valign
    first = True
    for prs in runs:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        _spacing(p, 0, 1)
        if align is not None:
            p.alignment = align
        for item in prs:
            t, b = item[0], item[1]
            col = item[2] if len(item) > 2 else None
            it = item[3] if len(item) > 3 else False
            _run(p, t, BODY_PT, b, col, it)
    return cell
