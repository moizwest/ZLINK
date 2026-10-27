"""Z-Link official document template — PERMANENT.

Artwork extracted from the genuine documents MW/10/1228 (tax invoice) and
ZL/PI/10356/01 (proforma). Applies the approved official design:
  header   navy bilingual banner, FULL BLEED, 0.8 mm from the top edge, every page
  ruler    A4, sides 10 mm, top 34 mm, bottom 19.8 mm, footer 5.3 mm
  watermark Z-Link logo at 81% page width, centred, vertical centre 52.5%,
            faded to 16% over white, behind all content, every page
  footer   document reference + Page X of Y in small grey, over the strapline
            artwork at 64% page width, centred, every page
  fills    NONE between the margins, so the watermark reads unbroken

Every Z-Link document builder imports this. Format changes are made here once.
"""
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement, parse_xml

import os
ART = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'artwork')
BANNER = f'{ART}/banner.png'
STRAPLINE = f'{ART}/strapline.png'
WATERMARK = f'{ART}/watermark.png'

NAVY = RGBColor(0x1F, 0x38, 0x64)
RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x80, 0x80, 0x80)
FONT = 'Calibri'

PAGE_W, PAGE_H = 210.0, 297.0          # mm, A4
MARGIN_SIDE, MARGIN_TOP = 10.0, 34.0
MARGIN_BOTTOM, FOOTER_DIST = 19.8, 5.3
HEADER_DIST = 0.8
USABLE = PAGE_W - 2 * MARGIN_SIDE      # 190.0 mm text block

# Geometry measured off the genuine proforma ZL/PI/10356/01 at 200 dpi.
TOP_BOXES = [94.9, 95.0]                               # equal halves
ITEM_COLS = [12.2, 86.4, 15.9, 14.2, 31.7, 29.5]       # SR/DESC/UNIT/QTY/RATE/TOTAL
BANK_COLS = [37.6, 80.0]                               # 117.6 mm overall, left aligned
assert abs(sum(ITEM_COLS) - USABLE) < 0.2, sum(ITEM_COLS)

BANNER_W = PAGE_W                      # full bleed
BANNER_H = 28.96
# Watermark: rendered from the ORIGINAL VECTOR LOGO, artwork/logo-master.pdf,
# supplied by the MD on 26 Oct 2026. Render the PDF at 1200 dpi, trim to the
# artwork, scale to WM_W and fade with factor 0.157 over white. Never rebuild it
# from a scan or by upscaling the small logo in the banner - both give a dirty,
# speckled, desaturated result. The vector logo's aspect is 0.900, matching the
# 0.898 footprint measured on the genuine document.
WM_X, WM_Y = 19.8, 79.2
WM_W, WM_H = 170.2, 153.1
WM_FADE = 0.157
STRAP_W = 139.4                        # ink measures 64.2% of page width
EMU_MM = 36000


def _float_behind(run, x_mm, y_mm, name):
    """Convert the run's inline picture into a page-anchored floating image
    placed behind the text, which is how a Word watermark is built."""
    inline = run._r.find(qn('w:drawing'))[0]
    extent = inline.find(qn('wp:extent'))
    cx, cy = extent.get('cx'), extent.get('cy')
    graphic = inline.find(qn('a:graphic'))
    docpr = inline.find(qn('wp:docPr'))
    docpr.set('name', name)

    anchor = parse_xml(
        '<wp:anchor %s distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="2" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>%d</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>%d</wp:posOffset></wp:positionV>'
        '<wp:extent cx="%s" cy="%s"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapNone/>'
        '</wp:anchor>' % (
            'xmlns:wp="%s"' % nsmap['wp'],
            int(x_mm * EMU_MM), int(y_mm * EMU_MM), cx, cy))
    anchor.append(docpr)
    cnv = inline.find(qn('wp:cNvGraphicFramePr'))
    if cnv is not None:
        anchor.append(cnv)
    anchor.append(graphic)
    drawing = run._r.find(qn('w:drawing'))
    drawing.remove(inline)
    drawing.append(anchor)


def new_document(doc_ref, show_page_numbers=True):
    """Return a Document already carrying the official Z-Link design."""
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(PAGE_W), Mm(PAGE_H)
    s.left_margin = s.right_margin = Mm(MARGIN_SIDE)
    s.top_margin = Mm(MARGIN_TOP)
    s.bottom_margin = Mm(MARGIN_BOTTOM)
    s.header_distance = Mm(HEADER_DIST)
    s.footer_distance = Mm(FOOTER_DIST)

    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    # ---- header: banner, full bleed, anchored to the page ----
    hp = s.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.space_before = Pt(0)
    hr = hp.add_run()
    hr.add_picture(BANNER, width=Mm(BANNER_W), height=Mm(BANNER_H))
    _float_behind(hr, 0.0, HEADER_DIST, 'Z-Link banner')

    # ---- watermark: anchored to the page, behind all content ----
    wr = hp.add_run()
    wr.add_picture(WATERMARK, width=Mm(WM_W), height=Mm(WM_H))
    _float_behind(wr, WM_X, WM_Y, 'Z-Link watermark')

    # ---- footer: reference line over the strapline artwork ----
    ref = s.footer.paragraphs[0]
    ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ref.paragraph_format.space_after = Pt(1)
    ref.paragraph_format.space_before = Pt(0)
    if show_page_numbers:
        def small(t):
            r = ref.add_run(t)
            r.font.size = Pt(7); r.font.color.rgb = GREY; r.font.name = FONT

        def field(instr):
            r = ref.add_run()
            b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
            i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = instr
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
            r._r.append(b); r._r.append(i); r._r.append(e)
            r.font.size = Pt(7); r.font.color.rgb = GREY; r.font.name = FONT

        small(f'{doc_ref}   |   Z-Link Technical Services Contracting Co. L.L.C.   |   Page ')
        field(' PAGE ')
        small(' of ')
        field(' NUMPAGES ')

    sp = s.footer.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.space_before = Pt(0)
    sp.add_run().add_picture(STRAPLINE, width=Mm(STRAP_W))

    return doc


# ---------- shared body helpers (no solid fills anywhere) ----------

def para(doc, text='', size=10, bold=False, color=None, align=None, after=2,
         italic=False, underline=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.underline = underline; r.font.name = FONT
    if color:
        r.font.color.rgb = color
    return p


def _twips(mm):
    return str(int(mm * 56.7))


def fix_grid(table, widths_mm, cell_margin_mm=1.2):
    """Pin the table to a fixed layout with explicit column widths.

    cell_margin_mm sets the left/right cell padding. The genuine documents sit
    text close to the rule; Word's 1.9 mm default makes narrow columns wrap.
    """
    tbl = table._tbl
    tblPr = tbl.tblPr
    for el in tblPr.findall(qn('w:tblCellMar')):
        tblPr.remove(el)
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('left', cell_margin_mm), ('right', cell_margin_mm),
                      ('top', 0.5), ('bottom', 0.5)):
        e = OxmlElement('w:%s' % side)
        e.set(qn('w:w'), str(int(val * 56.7)))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblPr.append(mar)
    for el in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(el)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    for g in tbl.findall(qn('w:tblGrid')):
        tbl.remove(g)
    grid = OxmlElement('w:tblGrid')
    for w in widths_mm:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), _twips(w)); grid.append(gc)
    tbl.insert(1, grid)
    for row in table.rows:
        for i, c in enumerate(row.cells):
            if i < len(widths_mm):
                c.width = Mm(widths_mm[i])


def cellset(cell, txt, size=10, bold=False, color=None, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    if align:
        p.alignment = align
    r = p.add_run(txt)
    r.font.size = Pt(size); r.bold = bold; r.font.name = FONT
    if color:
        r.font.color.rgb = color
    return cell


def cell_lines(cell, lines, size=10, first_bold=True, all_bold=False):
    cell.text = ''
    for k, ln in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(ln)
        r.font.size = Pt(size); r.font.name = FONT
        r.bold = all_bold or (first_bold and k == 0)
    return cell


def top_boxes(doc, client_lines, invoice_lines):
    """Clause B1 — the two side-by-side top boxes of a tax invoice / proforma.

    LEFT: the `TO,` block (client legal name, address, TRN, CONTACT, MOBILE,
    EMAIL ID). RIGHT: the invoice block (INVOICE/PROFORMA INVOICE no., Date,
    Contact: Moiz Husain, LPO No, Quote Ref, Project Location, Subject).

    MD RULES 27 Oct 2026, permanent — do not build these boxes any other way:
      * EVERY line in BOTH boxes is BOLD.
      * One BLANK line sits between the client TRN line and the CONTACT line
        in the left box (inserted here automatically if the caller omits it).
    """
    client_lines = list(client_lines)
    for i, ln in enumerate(client_lines[:-1]):
        if ln.strip().upper().startswith('TRN') and client_lines[i + 1].strip():
            client_lines.insert(i + 1, '')
            break
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    fix_grid(t, TOP_BOXES)
    cell_lines(t.cell(0, 0), client_lines, all_bold=True)
    cell_lines(t.cell(0, 1), invoice_lines, all_bold=True)
    return t


def align_items(table):
    """Clauses B3/B4 — alignment of the line-item table data (MD rule 27 Oct 2026).

    Every data row from the first line item down to NET TOTAL:
      * UNIT, QTY and RATE values: horizontally CENTRED, vertically MIDDLE.
      * TOTAL AMOUNT values: RIGHT-aligned, vertically MIDDLE.
      * SR NO: centred, middle.  DESCRIPTION: left, middle.
    The header row stays centred. Call after fix_grid() on the item table.
    Permanent — do not lay out the item table any other way.
    """
    C, R = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT
    col_align = {0: C, 2: C, 3: C, 4: C, 5: R}
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if row is table.rows[0]:
            continue
        seen = set()
        for i, cell in enumerate(row.cells):
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            txt = cell.text.strip().upper()
            if i == 4 and txt in ('GROSS TOTAL', 'VAT @ 5%', 'NET TOTAL'):
                continue            # totals labels stay right-aligned (B4)
            if txt.startswith('TOTAL IN WORDS'):
                continue
            if i in col_align:
                for p in cell.paragraphs:
                    p.alignment = col_align[i]


def words_row(table, amount_words):
    """Clause B5 — Total in Words as the LAST ROW of the line-item table
    (MD rule 28 Oct 2026, measured on genuine invoice MW/10/1228).

    The row is merged across all six columns, bold, black, directly beneath
    NET TOTAL. Never place Total in Words as a paragraph below the table.
    Call after the GROSS TOTAL / VAT / NET TOTAL rows and before align_items().
    """
    row = table.add_row()
    c = row.cells[0]
    for i in range(1, len(row.cells)):
        c = c.merge(row.cells[i])
    cellset(c, 'Total in Words:  ' + amount_words, 10, True)
    return row


def section_heading(doc, text):
    """Clause A6 — navy bold heading with a thin navy rule beneath (no fill).
    Used for the numbered sections of reports and work completion reports."""
    p = para(doc, text, 11, True, color=NAVY, after=3)
    p.paragraph_format.space_before = Pt(8)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', '1F3864')):
        bt.set(qn(k), v)
    bd.append(bt); pPr.append(bd)
    return p


def kv_table(doc, rows, widths_mm):
    """Particulars / client tables for reports (clauses W2, W3, W7):
    label cells bold, value cells plain, borders only."""
    t = doc.add_table(rows=len(rows), cols=len(widths_mm))
    t.style = 'Table Grid'
    fix_grid(t, widths_mm)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cellset(t.cell(i, j), v, 10, bold=(j % 2 == 0))
    return t
